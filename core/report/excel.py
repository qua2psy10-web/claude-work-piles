"""設計計算書(Excel)の生成。

Markdown 版と同じ内容を、シート分けした Excel ブックとして出力する。
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.analysis.stability import StabilityReport
from core.models.project import DesignProject
from core.soil.liquefaction import LiquefactionAssessment
from core.standards import NF_SAFETY_FACTOR, SAFETY_FACTORS

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
NG_FILL = PatternFill("solid", fgColor="FFC7CE")
TITLE_FONT = Font(bold=True, size=12)


def _write_title(ws: Worksheet, row: int, text: str) -> int:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    return row + 2


def _write_table(
    ws: Worksheet, row: int, header: list[str], rows: list[list], judge_col: int | None = None
) -> int:
    for j, name in enumerate(header, start=1):
        cell = ws.cell(row=row, column=j, value=name)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    row += 1
    for data in rows:
        for j, value in enumerate(data, start=1):
            ws.cell(row=row, column=j, value=value)
        if judge_col is not None and str(data[judge_col - 1]) == "NG":
            for j in range(1, len(header) + 1):
                ws.cell(row=row, column=j).fill = NG_FILL
        row += 1
    return row + 1


def _autosize(ws: Worksheet, max_width: int = 40) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = max(len(line) for line in str(cell.value).split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 8), length + 2)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(width, max_width)


def build_workbook(
    project: DesignProject,
    report: StabilityReport | None = None,
    liquefaction: LiquefactionAssessment | None = None,
) -> bytes:
    """計算書 Excel ブックをバイト列で返す。"""
    wb = Workbook()
    _sheet_conditions(wb.active, project)
    if liquefaction is not None:
        _sheet_liquefaction(wb.create_sheet("液状化判定"), liquefaction)
    if report is not None:
        _sheet_bearing(wb.create_sheet("支持力"), report)
        for i, case in enumerate(report.cases, start=1):
            name = f"{i}_{case.loads.case.value}"[:31]
            _sheet_case(wb.create_sheet(name), report, case)
        if report.negative_friction is not None:
            _sheet_nf(wb.create_sheet("負の周面摩擦力"), report)
        _sheet_summary(wb.create_sheet("総括"), report)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sheet_conditions(ws: Worksheet, project: DesignProject) -> None:
    ws.title = "設計条件"
    row = _write_title(ws, 1, f"杭基礎設計計算書 — {project.name}")
    row = _write_title(ws, row, "1.1 地層構成")
    rows = []
    top = 0.0
    for layer in project.soil_profile.layers:
        bottom = top + layer.thickness
        rows.append(
            [
                layer.name,
                layer.soil_type.value,
                round(top, 2),
                round(bottom, 2),
                layer.thickness,
                layer.n_value,
                layer.gamma_t,
                layer.gamma_sat,
                layer.cohesion,
            ]
        )
        top = bottom
    row = _write_table(
        ws,
        row,
        ["層名", "土質", "上端(m)", "下端(m)", "層厚(m)", "N値", "γt", "γsat", "c"],
        rows,
    )
    row = _write_table(
        ws,
        row,
        ["項目", "値"],
        [
            ["地下水位 (m)", project.soil_profile.gwl],
            ["地盤種別", project.seismic.ground_type.value],
            ["cIz", project.seismic.cz_type1],
            ["cIIz", project.seismic.cz_type2],
        ],
    )
    if project.pile is not None:
        rows = [
            ["杭種", project.pile.pile_type.value],
            ["施工工法", project.pile.method.value],
            ["杭径 D (m)", project.pile.diameter],
            ["杭長 L (m)", project.pile.length],
        ]
        if project.pile.wall_thickness is not None:
            rows.append(["板厚 t (mm)", project.pile.wall_thickness])
        if project.arrangement is not None:
            rows.append(["杭本数", project.arrangement.total_piles])
            rows.append(["杭間隔 (m)", project.arrangement.spacing_x])
        if project.footing is not None:
            rows.append(["フーチング厚 (m)", project.footing.height])
            rows.append(["根入れ深さ (m)", project.footing.embedment])
        row = _write_title(ws, row, "1.2 杭諸元")
        row = _write_table(ws, row, ["項目", "値"], rows)
    if project.loads:
        row = _write_title(ws, row, "1.3 荷重")
        row = _write_table(
            ws,
            row,
            ["荷重ケース", "V (kN)", "H (kN)", "M (kN·m)"],
            [[load.case.value, load.v, load.h, load.m] for load in project.loads],
        )
    _autosize(ws)


def _sheet_liquefaction(ws: Worksheet, assessment: LiquefactionAssessment) -> None:
    row = _write_title(ws, 1, "液状化の判定(道示Ⅴ 8.2)")
    row = _write_table(
        ws,
        row,
        ["地震動タイプ", "判定"],
        [
            ["タイプI", "液状化あり" if assessment.liquefiable_type1 else "液状化なし"],
            ["タイプII", "液状化あり" if assessment.liquefiable_type2 else "液状化なし"],
        ],
    )
    rows = []
    for s in assessment.slices:
        rows.append(
            [
                round(s.depth_top, 2),
                round(s.depth_bottom, 2),
                s.layer_name,
                s.soil_type.value,
                "○" if s.is_target else "対象外",
                s.excluded_reason or "",
                None if s.sigma_v is None else round(s.sigma_v, 1),
                None if s.sigma_v_eff is None else round(s.sigma_v_eff, 1),
                None if s.na is None else round(s.na, 2),
                None if s.rl is None else round(s.rl, 4),
                None if s.fl_type1 is None else round(s.fl_type1, 3),
                round(s.de_type1, 3),
                None if s.fl_type2 is None else round(s.fl_type2, 3),
                round(s.de_type2, 3),
            ]
        )
    _write_table(
        ws,
        row,
        [
            "上端(m)", "下端(m)", "層名", "土質", "対象", "対象外理由",
            "σv", "σ'v", "Na", "RL", "FL(I)", "DE(I)", "FL(II)", "DE(II)",
        ],
        rows,
    )
    _autosize(ws)


def _sheet_bearing(ws: Worksheet, report: StabilityReport) -> None:
    bc = report.bearing
    row = _write_title(ws, 1, "杭の軸方向支持力(道示Ⅳ 12.4)")
    row = _write_table(
        ws,
        row,
        ["項目", "値", "単位"],
        [
            ["先端支持力度 qd", round(bc.qd, 1), "kN/m²"],
            ["先端面積 Ap", round(bc.tip_area, 4), "m²"],
            ["先端支持力 qd·Ap", round(bc.tip_resistance, 1), "kN"],
            ["周面摩擦力 U·ΣLf", round(bc.skin_resistance, 1), "kN"],
            ["極限支持力 Ru", round(bc.ru, 1), "kN"],
            ["杭の有効重量 W", round(bc.w_pile, 1), "kN"],
            ["置換土の有効重量 Ws", round(bc.w_soil, 1), "kN"],
        ],
    )
    row = _write_table(
        ws,
        row,
        ["層名", "土質", "長さ(m)", "f (kN/m²)", "U·L·f (kN)"],
        [
            [s.layer_name, s.soil_type.value, round(s.length, 2), round(s.f, 1), round(s.force, 1)]
            for s in bc.skin_segments
        ],
    )
    _write_table(
        ws,
        row,
        ["荷重ケース", "n(押込み)", "Ra (kN)", "n(引抜き)", "Pa (kN)"],
        [
            [
                case.loads.case.value,
                SAFETY_FACTORS[case.loads.case.value][0],
                round(bc.allowable_push(case.loads.case), 1),
                SAFETY_FACTORS[case.loads.case.value][1],
                round(bc.allowable_pull(case.loads.case), 1),
            ]
            for case in report.cases
        ],
    )
    _autosize(ws)


def _sheet_case(ws: Worksheet, report: StabilityReport, case) -> None:
    sp = case.springs
    row = _write_title(ws, 1, f"安定計算 — {case.loads.case.value}")
    row = _write_table(
        ws,
        row,
        ["項目", "値", "単位"],
        [
            ["変形係数 E0", round(sp.e0, 0), "kN/m²"],
            ["換算載荷幅 BH", round(sp.bh, 4), "m"],
            ["水平方向地盤反力係数 kH", round(sp.kh, 0), "kN/m³"],
            ["特性値 β", round(sp.beta, 5), "1/m"],
            ["βL", round(sp.beta_le, 2), "—"],
            ["軸方向バネ Kv", round(case.kv, 0), "kN/m"],
            ["K1", round(sp.k1, 0), "kN/m"],
            ["K2 = K3", round(sp.k2, 0), "kN/rad"],
            ["K4", round(sp.k4, 0), "kN·m/rad"],
            ["水平変位 δ", round(case.result.u * 1000, 3), "mm"],
            ["回転角 θ", case.result.theta, "rad"],
        ],
    )
    row = _write_table(
        ws,
        row,
        ["杭", "x (m)", "軸力 (kN)", "水平力 (kN)", "杭頭M (kN·m)"],
        [
            [r.index, round(r.x, 3), round(r.axial, 1), round(r.shear, 1), round(r.moment, 1)]
            for r in case.result.reactions
        ],
    )
    row = _write_table(
        ws,
        row,
        ["照査項目", "作用値", "制限値", "単位", "比", "判定"],
        [
            [c.name, round(c.demand, 5), round(c.capacity, 5), c.unit, round(c.ratio, 3), c.judgement]
            for c in case.checks
        ],
        judge_col=6,
    )
    stress_rows = []
    for label, stress in (("杭頭", case.stress_head), ("地中部最大曲げ", case.stress_max)):
        if stress is None:
            continue
        for c in stress.checks:
            stress_rows.append(
                [
                    label,
                    round(stress.depth, 2),
                    c.name,
                    round(c.stress, 3),
                    round(c.allowable, 3),
                    round(c.ratio, 3),
                    c.judgement,
                ]
            )
    if stress_rows:
        row = _write_table(
            ws,
            row,
            ["位置", "深さ(m)", "照査項目", "応力度 (N/mm²)", "許容値", "比", "判定"],
            stress_rows,
            judge_col=7,
        )
    if case.pile_head is not None:
        row = _write_table(
            ws,
            row,
            ["杭頭結合部 照査項目", "応力度 (N/mm²)", "許容値", "比", "判定"],
            [
                [c.name, round(c.stress, 4), round(c.allowable, 3), round(c.ratio, 3), c.judgement]
                for c in case.pile_head.checks
            ],
            judge_col=5,
        )
    if case.forces is not None:
        _write_table(
            ws,
            row,
            ["深さ (m)", "変位 (mm)", "曲げM (kN·m)", "せん断力 (kN)"],
            [
                [round(p.depth, 3), round(p.displacement * 1000, 3), round(p.moment, 2), round(p.shear, 2)]
                for p in case.forces.points
            ],
        )
    _autosize(ws)


def _sheet_nf(ws: Worksheet, report: StabilityReport) -> None:
    nf = report.negative_friction
    row = _write_title(ws, 1, "負の周面摩擦力の検討(道示Ⅳ 12.4.3)")
    row = _write_table(
        ws,
        row,
        ["層名", "上端(m)", "下端(m)", "fn (kN/m²)", "NF (kN)"],
        [
            [s.layer_name, round(s.depth_top, 2), round(s.depth_bottom, 2), round(s.fn, 1), round(s.force, 1)]
            for s in nf.segments
        ],
    )
    _write_table(
        ws,
        row,
        ["項目", "値", "単位", "判定"],
        [
            ["中立点深さ", round(nf.neutral_depth, 2), "m", ""],
            ["負の周面摩擦力 NF", round(nf.nf, 1), "kN", ""],
            ["死荷重軸力", round(nf.dead_load, 1), "kN", ""],
            ["最大軸力 Nmax", round(nf.n_max, 1), "kN", ""],
            [f"許容値 Ru/{NF_SAFETY_FACTOR}", round(nf.allowable, 1), "kN", nf.judgement],
        ],
        judge_col=4,
    )
    _autosize(ws)


def _sheet_summary(ws: Worksheet, report: StabilityReport) -> None:
    row = _write_title(ws, 1, "総括")
    rows = [[case.loads.case.value, "OK" if case.all_ok else "NG"] for case in report.cases]
    if report.negative_friction is not None:
        rows.append(["負の周面摩擦力", report.negative_friction.judgement])
    rows.append(["総合判定", "OK" if report.all_ok else "NG"])
    _write_table(ws, row, ["項目", "判定"], rows, judge_col=2)
    _autosize(ws)
