"""杭体のせん断照査(道示Ⅳ(H24) 5.1.3)のテスト。

期待値は原典(スキャン)の規定から手計算で導いている。
docs/VERIFICATION.md 第23回を参照。
"""
import math

import pytest

from core.models import ConstructionMethod, LoadCase, PileSpec, PileType
from core.section.rc import RebarLayout
from core.section.shear import (
    ce_factor,
    check_shear,
    cn_factor,
    cpt_factor,
    effective_depth,
    equivalent_square_width,
    tensile_rebar_ratio,
    tension_quarter_positions,
)
from core.standards import TAU_A1_CONCRETE, TAU_A2_CONCRETE, TAU_C_CONCRETE

CIP = PileSpec(
    pile_type=PileType.CAST_IN_PLACE,
    method=ConstructionMethod.CAST_IN_PLACE,
    diameter=1.0,
    length=20.0,
)
REBAR = RebarLayout(count=24, diameter_mm=25.0, cover_mm=125.0)


# --- 断面諸元(図-解4.2.2) -------------------------------------------------


def test_equivalent_square_width():
    """b は円形断面と面積の等しい正方形の幅。"""
    b = equivalent_square_width(1.0)
    assert b**2 == pytest.approx(math.pi * 1.0**2 / 4.0)
    assert b == pytest.approx(0.8862269, abs=1e-6)
    # 幅は杭径より小さい(面積が等しい正方形なので)
    assert b < 1.0


def test_tension_quarter_selects_the_bottom_90_degrees():
    """引張側 1/4 部分は引張縁方向から ±45°。24本なら7本が該当する。"""
    quarter = tension_quarter_positions(1.0, REBAR)
    assert len(quarter) == 7
    r = REBAR.radius(1.0)
    # すべて引張側で、±45° の境界より外側にはない
    assert all(y <= -r * math.cos(math.pi / 4) + 1e-9 for y in quarter)
    # 最下端の鉄筋(y = −r)が含まれる
    assert min(quarter) == pytest.approx(-r)


def test_effective_depth_hand_calculation():
    """d = b/2 −(引張側 1/4 部分の鉄筋の重心)。

    D=1.0m、24-D25、かぶり125mm:
      r = 0.375、該当は 135°〜225° の 7 本
      重心 y = 0.375・(cos135+cos150+…+cos225)/7 = −0.325613
      b/2 = 0.443113
      d = 0.443113 + 0.325613 = 0.768726 m
    """
    r = 0.375
    angles = [math.radians(a) for a in (135, 150, 165, 180, 195, 210, 225)]
    centroid = sum(r * math.cos(a) for a in angles) / len(angles)
    expected = equivalent_square_width(1.0) / 2.0 - centroid

    assert effective_depth(1.0, REBAR) == pytest.approx(expected)
    assert effective_depth(1.0, REBAR) == pytest.approx(0.768726, abs=1e-6)
    # 有効高は杭径より小さく、かつ半分より大きい
    assert 0.5 < effective_depth(1.0, REBAR) < 1.0


def test_effective_depth_requires_bars_in_the_quarter():
    sparse = RebarLayout(count=4, diameter_mm=25.0, cover_mm=125.0)
    # 4本(0/90/180/270°)なら 180° の1本が該当する
    assert len(tension_quarter_positions(1.0, sparse)) == 1
    assert effective_depth(1.0, sparse) > 0


def test_tensile_rebar_ratio_excludes_bars_on_the_centroid():
    """図心位置(y=0)にちょうど乗る鉄筋は引張側に数えない。

    24本では 90°・270° の2本が y = 0 に乗る。浮動小数の丸めで符号が
    揺れるため、明示的に除外して決定的にしている(安全側)。
    """
    b = equivalent_square_width(1.0)
    d = effective_depth(1.0, REBAR)
    pt = tensile_rebar_ratio(1.0, REBAR, b, d)
    # 24本のうち引張側は 11 本(90°・270° は除く)
    expected = 100.0 * 11 * REBAR.bar_area / (b * d)
    assert pt == pytest.approx(expected)
    # 12本と数えた場合とは異なる(そちらは非安全側)
    assert pt < 100.0 * 12 * REBAR.bar_area / (b * d)


# --- 補正係数(表-4.2.2、表-4.2.3、式4.2.1) --------------------------------


def test_ce_factor_table_and_interpolation():
    assert ce_factor(0.300) == pytest.approx(1.4)
    assert ce_factor(1.000) == pytest.approx(1.0)
    assert ce_factor(3.000) == pytest.approx(0.7)
    assert ce_factor(10.000) == pytest.approx(0.5)
    # 範囲外は端の値で頭打ち
    assert ce_factor(0.100) == pytest.approx(1.4)
    assert ce_factor(20.000) == pytest.approx(0.5)
    # 線形補間(650mm は 300 と 1000 の中点 → 1.4 と 1.0 の中点)
    assert ce_factor(0.650) == pytest.approx(1.2)


def test_cpt_factor_table_and_interpolation():
    assert cpt_factor(0.1) == pytest.approx(0.7)
    assert cpt_factor(0.3) == pytest.approx(1.0)
    assert cpt_factor(1.0) == pytest.approx(1.5)
    assert cpt_factor(0.05) == pytest.approx(0.7)  # 頭打ち
    assert cpt_factor(2.0) == pytest.approx(1.5)  # 頭打ち
    # 0.15 は 0.1 と 0.2 の中点 → 0.7 と 0.9 の中点
    assert cpt_factor(0.15) == pytest.approx(0.8)


def test_cn_factor_formula_and_clamping():
    """cN = 1 + M0/M、M0 =(N/Ac)(Ic/y)、1 ≤ cN ≤ 2。"""
    d = 1.0
    area = math.pi * d**2 / 4.0
    section_modulus = math.pi * d**3 / 32.0
    axial, moment = 1000.0, 2000.0
    m0 = axial / area * section_modulus
    assert cn_factor(d, axial, moment) == pytest.approx(1.0 + m0 / moment)

    # 軸力が大きいと上限 2 で頭打ち
    assert cn_factor(d, 100000.0, 100.0) == 2.0
    # 引張軸力・M=0 は補正しない
    assert cn_factor(d, -500.0, 1000.0) == 1.0
    assert cn_factor(d, 1000.0, 0.0) == 1.0
    # 曲げが大きいほど補正は小さい(1 に近づく)
    assert cn_factor(d, 1000.0, 100000.0) < cn_factor(d, 1000.0, 1000.0)


# --- 照査 -------------------------------------------------------------------


def _result(case=LoadCase.PERMANENT, shear=400.0, moment=800.0, axial=1500.0):
    return check_shear(
        CIP, REBAR, 24, case, depth=0.0, shear=shear, moment=moment, axial=axial
    )


def test_tau_m_is_shear_over_b_times_d():
    """τm = Sh/(b・d)。単位は kN・m → N・mm。"""
    r = _result(shear=400.0)
    expected = 400.0 * 1000.0 / (r.width * 1000.0 * r.effective_depth * 1000.0)
    assert r.tau_m == pytest.approx(expected)
    # 参考: D=1.0m の杭で 400kN なら 0.6 N/mm² 程度
    assert r.tau_m == pytest.approx(0.5872, abs=1e-4)


def test_allowable_is_the_corrected_tau_a1():
    r = _result()
    assert r.tau_a1 == pytest.approx(
        r.ce * r.cpt * r.cn * TAU_A1_CONCRETE[24] * 1.0
    )
    assert r.tau_a2 == pytest.approx(TAU_A2_CONCRETE[24] * 1.0)


def test_storm_case_applies_the_increase_to_tau_a1():
    normal = _result(LoadCase.PERMANENT)
    storm = _result(LoadCase.STORM)
    assert storm.tau_a1 == pytest.approx(normal.tau_a1 * 1.25)
    assert storm.tau_a2 == pytest.approx(normal.tau_a2 * 1.25)
    # 作用応力度は変わらない
    assert storm.tau_m == pytest.approx(normal.tau_m)


def test_seismic_uses_tau_c_instead_of_tau_a1_times_1_5():
    """地震時は τa1×1.50 の代わりに τc を用いる(原典 4.2 の解説)。

    τc は τa1 を丸めた値の 1.5 倍と厳密には一致しないため、
    τa1×1.5 で代用すると値がずれる。
    """
    seismic = _result(LoadCase.LEVEL1_EQ)
    normal = _result(LoadCase.PERMANENT)
    base = seismic.tau_a1 / (seismic.ce * seismic.cpt * seismic.cn)
    assert base == pytest.approx(TAU_C_CONCRETE[24])
    assert base != pytest.approx(TAU_A1_CONCRETE[24] * 1.5)
    assert seismic.seismic and not normal.seismic
    # τa2 のほうは通常どおり割増しする
    assert seismic.tau_a2 == pytest.approx(TAU_A2_CONCRETE[24] * 1.5)


def test_stirrup_requirement_and_required_area():
    """τm > τa1 なら斜引張鉄筋が必要。必要量は式(5.1.3) による。"""
    small = _result(shear=100.0)
    assert not small.needs_stirrup
    assert small.required_stirrup_ratio(160.0) == 0.0

    large = _result(shear=900.0)
    assert large.needs_stirrup
    # Sca = τa1・b・d
    sca = (
        large.tau_a1 * large.width * 1000.0 * large.effective_depth * 1000.0 / 1000.0
    )
    assert large.concrete_shear_capacity == pytest.approx(sca)
    # Aw/s = 1.15・Sh'/(σsa・d・(sinθ+cosθ))、θ=90° なら sin+cos = 1
    excess = 900.0 - sca
    expected = 1.15 * excess * 1000.0 / (160.0 * large.effective_depth * 1000.0)
    assert large.required_stirrup_ratio(160.0) == pytest.approx(expected)
    # 45° に折り曲げると (sin+cos)=√2 なので必要量は減る
    assert large.required_stirrup_ratio(160.0, 45.0) == pytest.approx(
        expected / math.sqrt(2.0)
    )


def test_tau_a2_is_the_pass_fail_limit():
    """合否は τa2 で決まる(τa1 超過は「鉄筋が要る」であって NG ではない)。"""
    needs = _result(shear=900.0)
    assert needs.needs_stirrup
    assert needs.all_ok  # τa2 は超えていない

    over = _result(shear=3000.0)
    assert not over.all_ok
    assert over.tau_m > over.tau_a2


def test_non_cast_in_place_is_rejected():
    steel = PileSpec(
        pile_type=PileType.STEEL_PIPE,
        method=ConstructionMethod.DRIVEN,
        diameter=1.0,
        length=20.0,
        wall_thickness=12.0,
    )
    with pytest.raises(ValueError, match="未実装"):
        check_shear(steel, REBAR, 24, LoadCase.PERMANENT, 0.0, 400.0, 800.0, 1500.0)


def test_unsupported_fck_is_rejected():
    with pytest.raises(ValueError, match="許容せん断応力度"):
        check_shear(CIP, REBAR, 50, LoadCase.PERMANENT, 0.0, 400.0, 800.0, 1500.0)


def test_stability_runs_the_shear_check_at_the_max_shear_section():
    """安定計算がせん断力最大の断面で照査すること。"""
    from core.analysis.stability import analyze
    from core.models import (
        Footing,
        FootingLoads,
        PileArrangement,
        SoilLayer,
        SoilProfile,
        SoilType,
    )
    from core.section.checks import MaterialSpec

    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    report = analyze(
        CIP,
        PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        profile,
        [FootingLoads(case=LoadCase.LEVEL1_EQ, v=9000.0, h=3000.0, m=8000.0)],
        fck=24,
        material=MaterialSpec(fck=24, rebar=REBAR),
    )
    case = report.cases[0]
    assert case.shear is not None
    # 照査断面は断面力分布のせん断力最大点
    peak = case.forces.max_shear
    assert case.shear.depth == pytest.approx(peak.depth)
    assert case.shear.shear == pytest.approx(peak.shear)
    # 杭頭で最大になる(Chang の式)
    assert case.shear.depth == 0.0


def test_steel_pipe_pile_has_no_shear_result():
    """場所打ち杭以外は None(応力度照査側で未実装の注記が出る)。"""
    from core.analysis.stability import analyze
    from core.models import (
        Footing,
        FootingLoads,
        PileArrangement,
        SoilLayer,
        SoilProfile,
        SoilType,
    )
    from core.section.checks import MaterialSpec

    steel = PileSpec(
        pile_type=PileType.STEEL_PIPE,
        method=ConstructionMethod.DRIVEN,
        diameter=1.0,
        length=20.0,
        wall_thickness=12.0,
    )
    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    report = analyze(
        steel,
        PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        profile,
        [FootingLoads(case=LoadCase.PERMANENT, v=9000.0, h=600.0, m=1500.0)],
        fck=24,
        material=MaterialSpec(fck=24, rebar=REBAR),
    )
    assert report.cases[0].shear is None


# --- 計算書・GUI への反映 ----------------------------------------------------


def _cip_report(shear_load: float = 3000.0):
    from core.analysis.stability import analyze
    from core.models import (
        Footing,
        FootingLoads,
        PileArrangement,
        SoilLayer,
        SoilProfile,
        SoilType,
    )
    from core.section.checks import MaterialSpec

    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    from core.models import DesignProject

    project = DesignProject(
        pile=CIP,
        arrangement=PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        footing=Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        soil_profile=profile,
        loads=[
            FootingLoads(case=LoadCase.LEVEL1_EQ, v=9000.0, h=shear_load, m=8000.0)
        ],
    )
    report = analyze(
        project.pile, project.arrangement, project.footing,
        project.soil_profile,
        project.loads, fck=24, material=MaterialSpec(fck=24, rebar=REBAR),
    )
    return project, report


def test_markdown_report_includes_the_shear_section():
    from core.report.markdown import build_report

    project, report = _cip_report()
    text = build_report(project, report)
    assert "杭体のせん断照査" in text
    assert "換算幅 b" in text
    assert "図-解4.2.2" in text
    # 地震時は τc を用いた旨が出る
    assert "τc" in text


def test_markdown_report_states_the_required_stirrup():
    """斜引張鉄筋が必要なときは必要量まで出すこと。"""
    from core.report.markdown import build_report

    project, report = _cip_report(shear_load=12000.0)
    assert report.cases[0].shear.needs_stirrup
    text = build_report(project, report)
    assert "斜引張鉄筋が必要" in text
    assert "Aw/s" in text


def test_excel_report_includes_the_shear_table():
    import io

    import openpyxl

    from core.report.excel import build_workbook

    project, report = _cip_report()
    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(project, report)))
    ws = wb["1_レベル1地震時"]
    labels = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert any(v and "せん断照査" in str(v) for v in labels)
    assert any(v and "有効高 d" in str(v) for v in labels)


# --- 帯鉄筋(必要量 vs 配置量)------------------------------------------------

from core.section.rc import StirrupLayout  # noqa: E402
from core.section.shear import shear_capacity_level2  # noqa: E402

STIRRUP = StirrupLayout(diameter_mm=13.0, spacing_mm=150.0)


def test_stirrup_layout_geometry():
    """Aw = legs × 1本の断面積、Aw/s は間隔で割った値。"""
    assert STIRRUP.bar_area == pytest.approx(math.pi * 13.0**2 / 4.0)
    assert STIRRUP.area == pytest.approx(2 * STIRRUP.bar_area)
    assert STIRRUP.aw_per_spacing == pytest.approx(STIRRUP.area / 150.0)
    with pytest.raises(ValueError, match="径・間隔"):
        StirrupLayout(diameter_mm=0.0, spacing_mm=150.0).validated()
    with pytest.raises(ValueError, match="角度"):
        StirrupLayout(diameter_mm=13.0, spacing_mm=150.0, angle_deg=120.0).validated()


def test_stirrup_check_compares_required_against_provided():
    """帯鉄筋を与えると必要量との比較まで行うこと。"""
    # S = 450 kN なら D13@150(2本)で足りる
    result = check_shear(
        CIP, REBAR, 24, LoadCase.PERMANENT, 0.0, 450.0, 800.0, 1500.0,
        rebar_grade="SD345", stirrup=STIRRUP,
    )
    check = result.stirrup
    assert check is not None
    assert result.needs_stirrup  # τa1 は超えている
    assert check.provided == pytest.approx(STIRRUP.aw_per_spacing)
    assert check.required == pytest.approx(
        result.required_stirrup_ratio(check.sigma_sa)
    )
    # 場所打ち杭・常時なので σsa は水中部材の 160
    assert check.sigma_sa == pytest.approx(160.0)
    # この配筋なら足りている
    assert check.ok and result.all_ok


def test_insufficient_stirrup_makes_the_case_ng():
    """配置量が必要量に満たなければ NG になること(τa2 は満たしていても)。"""
    result = check_shear(
        CIP, REBAR, 24, LoadCase.PERMANENT, 0.0, 900.0, 800.0, 1500.0,
        rebar_grade="SD345", stirrup=STIRRUP,
    )
    assert result.stirrup is not None
    assert not result.stirrup.ok
    assert not result.all_ok
    # τa2 のほうは満たしている(NG の原因は鉄筋量)
    assert all(c.ok for c in result.checks)


def test_stirrup_check_requires_a_rebar_grade():
    with pytest.raises(ValueError, match="鉄筋材質"):
        check_shear(
            CIP, REBAR, 24, LoadCase.PERMANENT, 0.0, 900.0, 800.0, 1500.0,
            stirrup=STIRRUP,
        )


def test_seismic_stirrup_uses_the_other_category_not_axial_rebar():
    """斜引張鉄筋の地震時の基本値は「上記以外」の 200(軸方向鉄筋ではない)。"""
    result = check_shear(
        CIP, REBAR, 24, LoadCase.LEVEL1_EQ, 0.0, 900.0, 800.0, 1500.0,
        rebar_grade="SD490", stirrup=STIRRUP,
    )
    # SD490 でも斜引張鉄筋は 200 × 1.50 = 300(軸方向鉄筋なら 290×1.5 = 435)
    assert result.stirrup.sigma_sa == pytest.approx(300.0)


# --- レベル2のせん断耐力(道示Ⅳ 5.2.3)---------------------------------------


def test_shear_capacity_formula():
    """Ps = Sc + Ss。Sc = cc·ce·cpt·cN·τc·b·d、Ss = Aw·σsy·d·(sinθ+cosθ)/(1.15s)。"""
    cap = shear_capacity_level2(
        CIP, REBAR, 24, axial=1500.0, moment=2000.0, stirrup=STIRRUP
    )
    b_mm = cap.width * 1000.0
    d_mm = cap.effective_depth * 1000.0
    assert cap.cc == 1.0  # 基礎は cc = 1
    assert cap.sc == pytest.approx(
        cap.cc * cap.ce * cap.cpt * cap.cn * cap.tau_c * b_mm * d_mm / 1000.0
    )
    assert cap.ss == pytest.approx(
        STIRRUP.area * 345.0 * d_mm * 1.0 / (1.15 * STIRRUP.spacing_mm) / 1000.0
    )
    assert cap.total == pytest.approx(cap.sc + cap.ss)


def test_shear_capacity_caps_the_stirrup_yield_at_345():
    """斜引張鉄筋の降伏点は SD390・SD490 でも 345 で頭打ち(道示Ⅳ 5.2.3)。"""
    base = shear_capacity_level2(
        CIP, REBAR, 24, 1500.0, 2000.0, stirrup=STIRRUP, rebar_grade="SD345"
    )
    high = shear_capacity_level2(
        CIP, REBAR, 24, 1500.0, 2000.0, stirrup=STIRRUP, rebar_grade="SD490"
    )
    assert base.sigma_sy == 345.0
    assert high.sigma_sy == 345.0
    assert high.ss == pytest.approx(base.ss)


def test_shear_capacity_without_stirrup_is_concrete_only():
    cap = shear_capacity_level2(CIP, REBAR, 24, 1500.0, 2000.0)
    assert cap.ss == 0.0
    assert cap.sigma_sy is None
    assert cap.total == pytest.approx(cap.sc)


def test_shear_capacity_uses_tau_c_not_tau_a1():
    """レベル2は τc(表-5.2.1)を用いる。τa1 の 1.5 倍ではない。"""
    cap = shear_capacity_level2(CIP, REBAR, 24, 1500.0, 2000.0)
    assert cap.tau_c == pytest.approx(TAU_C_CONCRETE[24])
    assert cap.tau_c != pytest.approx(TAU_A1_CONCRETE[24] * 1.5)


def test_level2_reports_the_shear_capacity_check():
    """run_level2 がせん断耐力の照査を出すこと。"""
    from core.analysis.level2 import run_level2
    from core.models import Footing, PileArrangement, SoilLayer, SoilProfile, SoilType

    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    kwargs = dict(
        pile=CIP,
        arrangement=PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        footing=Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        profile=profile,
        v_load=9000.0, h_load=4000.0, m_load=12000.0,
        fck=24, rebar=REBAR, max_factor=1.0, steps=10,
    )
    with_stirrup = run_level2(**kwargs, stirrup=STIRRUP)
    without = run_level2(**kwargs)

    assert with_stirrup.shear_capacity is not None
    assert with_stirrup.response_shear is not None
    assert any("せん断耐力" in c.name for c in with_stirrup.checks)
    # 帯鉄筋があるぶん耐力が大きい
    assert with_stirrup.shear_capacity.total > without.shear_capacity.total
    assert without.shear_capacity.ss == 0.0
    assert any("帯鉄筋が未入力" in n for n in without.notes)


def test_level2_skips_the_shear_check_without_rebar():
    from core.analysis.level2 import run_level2
    from core.models import Footing, PileArrangement, SoilLayer, SoilProfile, SoilType

    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    result = run_level2(
        CIP,
        PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        profile,
        v_load=9000.0, h_load=4000.0, m_load=12000.0,
        fck=24, max_factor=1.0, steps=10,
    )
    assert result.shear_capacity is None
    assert any("軸方向鉄筋が未入力" in n for n in result.notes)


def test_reports_include_the_stirrup_and_level2_shear():
    """帯鉄筋の照査とレベル2のせん断耐力が計算書に出ること。"""
    from core.analysis.level2 import run_level2
    from core.analysis.stability import analyze
    from core.models import (
        DesignProject,
        Footing,
        FootingLoads,
        PileArrangement,
        SoilLayer,
        SoilProfile,
        SoilType,
    )
    from core.report.markdown import build_report
    from core.section.checks import MaterialSpec

    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    arrangement = PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5)
    footing = Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0)
    project = DesignProject(
        pile=CIP, arrangement=arrangement, footing=footing,
        soil_profile=profile,
        loads=[FootingLoads(case=LoadCase.LEVEL1_EQ, v=9000.0, h=3000.0, m=8000.0)],
    )
    report = analyze(
        CIP, arrangement, footing, profile, project.loads, fck=24,
        material=MaterialSpec(fck=24, rebar=REBAR, stirrup=STIRRUP),
    )
    level2 = run_level2(
        CIP, arrangement, footing, profile,
        v_load=9000.0, h_load=4000.0, m_load=12000.0,
        fck=24, rebar=REBAR, stirrup=STIRRUP, max_factor=1.0, steps=10,
    )
    text = build_report(project, report, level2=level2)
    assert "斜引張鉄筋量 Aw/s" in text
    assert "杭体のせん断耐力" in text
    assert "Ps = Sc + Ss" in text
    assert "345 N/mm² で頭打ち" in text
    # L1 側は必要量と配置量の両方が出る
    assert report.cases[0].shear.stirrup is not None


# --- 斜め圧縮破壊(道示Ⅲ 4.3.4、表-4.3.2)------------------------------------


def test_tau_max_table_pinned():
    """表-4.3.2 コンクリートの平均せん断応力度の最大値。"""
    from core import standards as st_mod

    assert st_mod.TAU_MAX_CONCRETE == {
        21: 2.8, 24: 3.2, 27: 3.6, 30: 4.0, 40: 5.3, 50: 6.0, 60: 6.0,
    }
    # τc(負担できる値)とは用途が異なり、桁が違う
    for fck, tau_c in st_mod.TAU_C_CONCRETE.items():
        assert st_mod.TAU_MAX_CONCRETE[fck] > tau_c * 5
    # 許容応力度法の τa2(斜引張鉄筋と共同の上限)より大きい(終局レベルのため)
    for fck, tau_a2 in st_mod.TAU_A2_CONCRETE.items():
        assert st_mod.TAU_MAX_CONCRETE[fck] > tau_a2


def test_web_crushing_capacity_formula():
    """Suc = τmax・bw・d(RC部材なので Sp = 0)。"""
    cap = shear_capacity_level2(
        CIP, REBAR, 24, axial=1500.0, moment=2000.0, stirrup=STIRRUP
    )
    assert cap.tau_max == pytest.approx(3.2)
    assert cap.web_crushing_capacity == pytest.approx(
        3.2 * cap.width * 1000.0 * cap.effective_depth * 1000.0 / 1000.0
    )


def test_web_crushing_caps_the_capacity():
    """斜引張鉄筋を増やしても Suc は超えられない。"""
    modest = shear_capacity_level2(
        CIP, REBAR, 24, 1500.0, 2000.0, stirrup=STIRRUP
    )
    heavy = shear_capacity_level2(
        CIP, REBAR, 24, 1500.0, 2000.0,
        stirrup=StirrupLayout(diameter_mm=22.0, spacing_mm=75.0),
    )
    # Sus は鉄筋量で伸びるが Suc は変わらない
    assert heavy.total > modest.total
    assert heavy.web_crushing_capacity == pytest.approx(
        modest.web_crushing_capacity
    )
    # 過密配筋では斜め圧縮破壊が支配する
    assert not modest.web_crushing_governs
    assert heavy.web_crushing_governs
    assert heavy.governing_capacity == pytest.approx(heavy.web_crushing_capacity)
    assert modest.governing_capacity == pytest.approx(modest.total)


def test_level2_checks_both_failure_modes():
    """レベル2は斜引張破壊と斜め圧縮破壊の両方を照査すること。"""
    from core.analysis.level2 import run_level2
    from core.models import Footing, PileArrangement, SoilLayer, SoilProfile, SoilType

    profile = SoilProfile(
        layers=[
            SoilLayer(name="As", soil_type=SoilType.SAND, thickness=10.0,
                      n_value=15.0, gamma_t=18.0, gamma_sat=19.0),
            SoilLayer(name="Ds", soil_type=SoilType.SAND, thickness=25.0,
                      n_value=45.0, gamma_t=19.0, gamma_sat=20.0),
        ],
        gwl=2.0,
    )
    result = run_level2(
        CIP,
        PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        profile,
        v_load=9000.0, h_load=4000.0, m_load=12000.0,
        fck=24, rebar=REBAR, stirrup=STIRRUP, max_factor=1.0, steps=10,
    )
    names = [c.name for c in result.checks]
    assert "杭体のせん断耐力(斜引張破壊)" in names
    assert "コンクリートの斜め圧縮破壊" in names
    # 同じ応答せん断力を2つの耐力と比べている
    shear_checks = [c for c in result.checks if "破壊" in c.name]
    assert len({round(c.demand, 6) for c in shear_checks}) == 1
