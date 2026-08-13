"""計算書生成のテスト。"""
import io

from openpyxl import load_workbook

from core.analysis.stability import analyze
from core.models import (
    ConstructionMethod,
    DesignProject,
    Footing,
    FootingLoads,
    LoadCase,
    PileArrangement,
    PileSpec,
    PileType,
    SeismicConditions,
    SoilLayer,
    SoilProfile,
    SoilType,
)
from core.report.excel import build_workbook
from core.report.markdown import build_report
from core.section.checks import MaterialSpec
from core.section.rc import RebarLayout
from core.soil.liquefaction import assess_liquefaction
from core.standards import GroundType

MATERIAL = MaterialSpec(
    fck=24, rebar=RebarLayout(count=24, diameter_mm=25.0, cover_mm=125.0)
)


def sample_project() -> DesignProject:
    profile = SoilProfile(
        layers=[
            SoilLayer(
                name="As1", soil_type=SoilType.SAND, thickness=8.0, n_value=10.0,
                gamma_t=18.0, gamma_sat=19.0, fc=5.0, d50=0.3, d10=0.08,
            ),
            SoilLayer(
                name="Ac1", soil_type=SoilType.CLAY, thickness=5.0, n_value=4.0,
                gamma_t=16.0, gamma_sat=16.5, cohesion=40.0,
            ),
            SoilLayer(
                name="Ds1", soil_type=SoilType.SAND, thickness=17.0, n_value=40.0,
                gamma_t=19.0, gamma_sat=20.0, fc=8.0, d50=0.5, d10=0.1,
                is_alluvial=False,
            ),
        ],
        gwl=1.5,
    )
    return DesignProject(
        name="テスト橋",
        soil_profile=profile,
        seismic=SeismicConditions(ground_type=GroundType.TYPE_II),
        pile=PileSpec(
            pile_type=PileType.CAST_IN_PLACE,
            method=ConstructionMethod.CAST_IN_PLACE,
            diameter=1.0,
            length=20.0,
        ),
        arrangement=PileArrangement(nx=2, ny=3, spacing_x=2.5, spacing_y=2.5),
        footing=Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        loads=[
            FootingLoads(case=LoadCase.PERMANENT, v=9000.0, h=300.0, m=1500.0),
            FootingLoads(case=LoadCase.LEVEL1_EQ, v=9000.0, h=2000.0, m=8000.0),
        ],
    )


def sample_report(project: DesignProject):
    return analyze(
        project.pile,
        project.arrangement,
        project.footing,
        project.soil_profile,
        project.loads,
        fck=24,
        material=MATERIAL,
        check_negative_friction=True,
    )


def test_markdown_report_contains_all_sections():
    project = sample_project()
    report = sample_report(project)
    liq = assess_liquefaction(project.soil_profile, GroundType.TYPE_II)
    text = build_report(project, report, liq)

    for heading in (
        "# 杭基礎設計計算書",
        "## 1. 設計条件",
        "## 2. 液状化の判定",
        "## 3. 杭の軸方向支持力",
        "## 4. 安定計算 — 常時",
        "## 4. 安定計算 — レベル1地震時",
        "## 5. 軸方向鉄筋量の照査",
        "## 6. 負の周面摩擦力の検討",
        "## 7. 総括",
    ):
        assert heading in text
    assert "テスト橋" in text
    assert "総合判定" in text
    # 各荷重ケースの応力度照査・杭頭結合部が出力される
    assert "杭体の応力度照査" in text
    assert "杭頭結合部の照査" in text


def test_markdown_report_without_analysis():
    """安定計算前でも設計条件だけの計算書が作れること。"""
    text = build_report(sample_project())
    assert "## 1. 設計条件" in text
    assert "## 3. 杭の軸方向支持力" not in text


def test_markdown_reports_judgement():
    project = sample_project()
    report = sample_report(project)
    text = build_report(project, report)
    expected = "OK" if report.all_ok else "NG"
    assert f"**総合判定: {expected}**" in text


def test_excel_workbook_sheets():
    project = sample_project()
    report = sample_report(project)
    liq = assess_liquefaction(project.soil_profile, GroundType.TYPE_II)
    data = build_workbook(project, report, liq)

    wb = load_workbook(io.BytesIO(data))
    assert "設計条件" in wb.sheetnames
    assert "液状化判定" in wb.sheetnames
    assert "支持力" in wb.sheetnames
    assert "負の周面摩擦力" in wb.sheetnames
    assert "総括" in wb.sheetnames
    # 荷重ケースごとのシート
    assert any("常時" in name for name in wb.sheetnames)
    assert any("レベル1地震時" in name for name in wb.sheetnames)


def test_excel_summary_contains_judgement():
    project = sample_project()
    report = sample_report(project)
    wb = load_workbook(io.BytesIO(build_workbook(project, report)))
    ws = wb["総括"]
    values = [
        (row[0].value, row[1].value)
        for row in ws.iter_rows(min_row=1, max_col=2)
        if row[0].value
    ]
    assert ("総合判定", "OK" if report.all_ok else "NG") in values


def test_excel_sheet_names_within_limit():
    project = sample_project()
    report = sample_report(project)
    wb = load_workbook(io.BytesIO(build_workbook(project, report)))
    assert all(len(name) <= 31 for name in wb.sheetnames)


def test_excel_without_analysis():
    data = build_workbook(sample_project())
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["設計条件"]


def test_level2_section_is_included():
    """レベル2の結果を渡すと計算書に第7章が追加されること。"""
    from core.analysis.level2 import AxialSpringModel, analyze_level2
    from core.models import PileArrangement

    arrangement = PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5)
    axial = AxialSpringModel(kv=5.0e5, push_limit=1500.0, pull_limit=1500.0)
    result = analyze_level2(
        arrangement, axial, k1=2.0e5, k2=-1.0e5, k4=1.0e5,
        v_load=9000.0, h_load=2000.0, m_load=8000.0,
        allowable_ductility=4.0, allowable_displacement=0.3,
    )
    project = sample_project()
    md = build_report(project, level2=result)

    assert "レベル2地震時の照査" in md
    assert "応答塑性率" in md
    assert "制限事項" in md
    # 制限事項が省略されずすべて出ていること
    from core.analysis.level2 import LIMITATIONS
    for limitation in LIMITATIONS:
        assert limitation in md


def test_level2_section_absent_when_not_run():
    md = build_report(sample_project())
    assert "レベル2地震時の照査" not in md


def test_level2_report_includes_soil_reaction_diagnosis():
    """pHU との突合診断が計算書に含まれること。"""
    from core.analysis.level2 import run_level2
    from core.models import (
        ConstructionMethod,
        Footing,
        PileArrangement,
        PileSpec,
        PileType,
        SoilLayer,
        SoilProfile,
        SoilType,
    )

    profile = SoilProfile(
        layers=[
            SoilLayer(
                name="As", soil_type=SoilType.SAND, thickness=30.0, n_value=20.0,
                gamma_t=18.0, gamma_sat=19.0, k_ep=0.05,
            )
        ],
        gwl=2.0,
    )
    pile = PileSpec(
        pile_type=PileType.STEEL_PIPE, method=ConstructionMethod.DRIVEN,
        diameter=1.0, length=20.0, wall_thickness=12.0,
    )
    result = run_level2(
        pile,
        PileArrangement(nx=3, ny=3, spacing_x=2.5, spacing_y=2.5),
        Footing(width_x=8.0, width_y=8.0, height=1.5, embedment=2.0),
        profile,
        v_load=9000.0, h_load=3000.0, m_load=12000.0,
    )
    md = build_report(sample_project(), level2=result)

    assert "pHU" in md
    assert result.soil_reaction is not None
    # KEP を極端に小さくしたので超過し、非安全側である旨が出る
    assert not result.soil_reaction.ok
    assert "非安全側" in md


def test_report_shows_the_liquefaction_reduction():
    """DE を反映した場合、計算書のバネ定数欄に低減係数が出ること。"""
    from tests.test_soil_reduction import (
        ARRANGEMENT,
        FOOTING,
        PILE,
        liquefiable_profile,
        sample_reduction,
    )
    from core.models import FootingLoads, LoadCase

    loads = [FootingLoads(case=LoadCase.LEVEL1_EQ, v=9000.0, h=2000.0, m=8000.0)]
    report = analyze(
        PILE, ARRANGEMENT, FOOTING, liquefiable_profile(), loads,
        reduction=sample_reduction(),
    )
    md = build_report(sample_project(), report)
    assert "液状化による低減係数 DE" in md
    assert "液状化による土質定数の低減" in md  # 注記


def test_report_omits_the_reduction_row_when_not_applied():
    from tests.test_soil_reduction import (
        ARRANGEMENT,
        FOOTING,
        PILE,
        liquefiable_profile,
    )
    from core.models import FootingLoads, LoadCase

    loads = [FootingLoads(case=LoadCase.LEVEL1_EQ, v=9000.0, h=2000.0, m=8000.0)]
    report = analyze(PILE, ARRANGEMENT, FOOTING, liquefiable_profile(), loads)
    assert "液状化による低減係数" not in build_report(sample_project(), report)


def test_report_shows_the_group_pile_factor_only_when_it_applies():
    """μ を乗じた場合だけ、バネ定数の表に μ の行を出す。"""
    from core.analysis.stability import analyze
    from core.models import Footing, PileArrangement
    from core.report.markdown import build_report
    from tests.test_validation import sample_inputs, sample_loads

    pile, wide, _, profile = sample_inputs()
    footing = Footing(width_x=12.0, width_y=12.0, height=1.5, embedment=2.0)
    project = DesignProject(
        soil_profile=profile, pile=pile, footing=footing, loads=sample_loads()
    )

    narrow = PileArrangement(nx=3, ny=3, spacing_x=2.0, spacing_y=2.0)
    text = build_report(
        project.model_copy(update={"arrangement": narrow}),
        analyze(pile, narrow, footing, profile, sample_loads()),
    )
    assert "群杭の補正係数 μ" in text
    assert "0.900" in text

    text = build_report(
        project.model_copy(update={"arrangement": wide}),
        analyze(pile, wide, footing, profile, sample_loads()),
    )
    assert "群杭の補正係数" not in text
